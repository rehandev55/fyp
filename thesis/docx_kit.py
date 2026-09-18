"""
Rendering primitives for the final year project report.

Every rule in this module comes from "Project Report template-3.0.docx" and
from the red-ink notes inside it. The notes take priority over the template's
own sample formatting wherever the two disagree; see SUBHEADING_BOLD below.

Page setup (measured from the template):
    A4, 21.0 x 29.7 cm
    left margin  3.8 cm, right / top / bottom 3.0 cm
    header and footer distance 1.27 cm

Typography:
    Times New Roman 12 pt, line spacing 1.5, justified, first line indented
    chapter heading  14 pt, bold, ALL CAPITALS, centred
    sub-heading      12 pt, ALL CAPITALS, not bold
    sub-sub-heading  12 pt, Initial capital only, bold

Pagination:
    front matter  roman numerals, centred, hidden on the cover page
    body onwards  arabic numerals restarting at 1 on the first page of
                  Chapter 1, with a running header naming the chapter
"""

import os

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import (WD_ALIGN_PARAGRAPH, WD_LINE_SPACING,
                            WD_TAB_ALIGNMENT, WD_TAB_LEADER)
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

# ── template-derived constants ───────────────────────────────────────────────
BODY_FONT = "Times New Roman"
CODE_FONT = "Consolas"
BODY_SIZE = Pt(12)
LINE_SPACING = 1.5

PAGE_W, PAGE_H = Cm(21.0), Cm(29.7)
MARGIN_LEFT = Cm(3.8)
MARGIN_RIGHT = Cm(3.0)
MARGIN_TOP = Cm(3.0)
MARGIN_BOTTOM = Cm(3.0)
HEADER_DIST = Cm(1.27)
FOOTER_DIST = Cm(1.27)

#: usable text width, used for the dot-leader tab stop in the contents lists
CONTENT_WIDTH = Inches(5.59)

FIRST_LINE_INDENT = Inches(0.5)

#: The red note in the template reads: "The main heading should be bold and all
#: letters should be capital, The subheading should also be in capital letters,
#: but it should not be bold."  The template's own body text shows bold
#: sub-headings, which contradicts its note. The note wins. Flip this to True if
#: the supervisor prefers the sample formatting instead.
SUBHEADING_BOLD = False

GREY = RGBColor(0x55, 0x55, 0x55)
HEADER_BG = "D9D9D9"
ZEBRA_BG = "F2F2F2"
CODE_BG = "F5F5F5"
NOTE_BG = "FBF6E9"
SHOT_BG = "FAFAFA"

FIG_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "figures")


# ── low-level XML helpers ────────────────────────────────────────────────────
def _el(tag: str):
    return OxmlElement(tag)


def shade(cell_or_para, fill: str) -> None:
    """Apply a solid background fill to a table cell or a paragraph."""
    shd = _el("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    if hasattr(cell_or_para, "_tc"):
        cell_or_para._tc.get_or_add_tcPr().append(shd)
    else:
        cell_or_para._p.get_or_add_pPr().append(shd)


def set_cell_margins(table, top=60, bottom=60, left=90, right=90) -> None:
    """Set uniform cell padding, in twentieths of a point."""
    tbl_pr = table._tbl.tblPr
    mar = _el("w:tblCellMar")
    for side, value in (("top", top), ("bottom", bottom),
                        ("left", left), ("right", right)):
        node = _el(f"w:{side}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        mar.append(node)
    tbl_pr.append(mar)


def repeat_header(row) -> None:
    """Mark a table row so it repeats on every page the table spans."""
    tr_pr = row._tr.get_or_add_trPr()
    header = _el("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def keep_with_next(paragraph) -> None:
    """Prevent a caption or heading being orphaned at the foot of a page."""
    paragraph.paragraph_format.keep_with_next = True


def add_field(paragraph, instruction: str, placeholder: str = "1"):
    """Insert a Word field such as PAGE or NUMPAGES."""
    run = paragraph.add_run()
    begin = _el("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = _el("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" {instruction} "
    separate = _el("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = _el("w:t")
    text.text = placeholder
    end = _el("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (begin, instr, separate, text, end):
        run._r.append(node)
    return run


_BOOKMARK_ID = [1000]


def bookmark(paragraph, name: str) -> None:
    """Name a paragraph so its final page number can be read back from Word."""
    _BOOKMARK_ID[0] += 1
    bid = str(_BOOKMARK_ID[0])
    start = _el("w:bookmarkStart")
    start.set(qn("w:id"), bid)
    start.set(qn("w:name"), name)
    end = _el("w:bookmarkEnd")
    end.set(qn("w:id"), bid)
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def page_numbering(section, fmt: str = "decimal", start: int | None = None,
                   restart: bool = False) -> None:
    """Set the numeral style, and optionally restart the count, for a section."""
    sect_pr = section._sectPr
    existing = sect_pr.find(qn("w:pgNumType"))
    if existing is not None:
        sect_pr.remove(existing)
    node = _el("w:pgNumType")
    node.set(qn("w:fmt"), fmt)
    if restart and start is not None:
        node.set(qn("w:start"), str(start))
    sect_pr.append(node)


# ── document and section construction ────────────────────────────────────────
def _style_page(section) -> None:
    section.page_width = PAGE_W
    section.page_height = PAGE_H
    section.left_margin = MARGIN_LEFT
    section.right_margin = MARGIN_RIGHT
    section.top_margin = MARGIN_TOP
    section.bottom_margin = MARGIN_BOTTOM
    section.header_distance = HEADER_DIST
    section.footer_distance = FOOTER_DIST


def new_document() -> Document:
    doc = Document()

    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = BODY_SIZE
    rpr = normal.element.get_or_add_rPr().get_or_add_rFonts()
    rpr.set(qn("w:eastAsia"), BODY_FONT)
    rpr.set(qn("w:cs"), BODY_FONT)
    pf = normal.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.space_after = Pt(0)
    pf.space_before = Pt(0)

    _style_page(doc.sections[0])
    return doc


def new_section(doc, header_text: str = "", numbering: str = "decimal",
                start: int | None = None, restart: bool = False,
                different_first: bool = True):
    """Start a new page-level section with its own header and page numbering."""
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    _style_page(section)
    section.different_first_page_header_footer = different_first

    for header in (section.header, section.first_page_header):
        header.is_linked_to_previous = False
        for paragraph in list(header.paragraphs)[1:]:
            paragraph._p.getparent().remove(paragraph._p)
        header.paragraphs[0].text = ""

    if header_text:
        # Running header: right aligned, bold italic, as in the template.
        para = section.header.paragraphs[0]
        para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = para.add_run(header_text)
        run.bold = True
        run.italic = True
        run.font.name = BODY_FONT
        run.font.size = Pt(12)

    for footer in (section.footer, section.first_page_footer):
        footer.is_linked_to_previous = False
        for paragraph in list(footer.paragraphs)[1:]:
            paragraph._p.getparent().remove(paragraph._p)
        para = footer.paragraphs[0]
        para.text = ""
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = add_field(para, "PAGE")
        run.font.name = BODY_FONT
        run.font.size = Pt(12)

    page_numbering(section, numbering, start, restart)
    return section


def front_matter_section(doc) -> None:
    """Configure the opening section: roman numerals, no number on the cover."""
    section = doc.sections[0]
    section.different_first_page_header_footer = True

    section.first_page_footer.is_linked_to_previous = False
    section.first_page_footer.paragraphs[0].text = ""
    section.first_page_header.is_linked_to_previous = False
    section.first_page_header.paragraphs[0].text = ""

    footer = section.footer
    footer.is_linked_to_previous = False
    para = footer.paragraphs[0]
    para.text = ""
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = add_field(para, "PAGE", "ii")
    run.font.name = BODY_FONT
    run.font.size = Pt(12)

    page_numbering(section, "lowerRoman", start=1, restart=True)


# ── text primitives ──────────────────────────────────────────────────────────
def para(doc, text="", size=12, bold=False, italic=False, align="justify",
         indent=False, space_after=6, space_before=0, colour=None,
         font=None, caps=False, line_spacing=LINE_SPACING):
    paragraph = doc.add_paragraph()
    paragraph.alignment = {
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
    }[align]
    pf = paragraph.paragraph_format
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(space_before)
    if line_spacing:
        pf.line_spacing = line_spacing
    if indent:
        pf.first_line_indent = FIRST_LINE_INDENT
    if text:
        run = paragraph.add_run(text.upper() if caps else text)
        run.bold = bold
        run.italic = italic
        run.font.size = Pt(size)
        run.font.name = font or BODY_FONT
        if colour is not None:
            run.font.color.rgb = colour
    return paragraph


def rich(doc, chunks, align="justify", size=12, indent=False, space_after=6):
    """A paragraph assembled from (text, bold, italic) triples."""
    paragraph = doc.add_paragraph()
    paragraph.alignment = {
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
    }[align]
    pf = paragraph.paragraph_format
    pf.space_after = Pt(space_after)
    pf.line_spacing = LINE_SPACING
    if indent:
        pf.first_line_indent = FIRST_LINE_INDENT
    for chunk in chunks:
        text, bold, italic = (list(chunk) + [False, False])[:3]
        run = paragraph.add_run(text)
        run.bold = bold
        run.italic = italic
        run.font.size = Pt(size)
        run.font.name = BODY_FONT
    return paragraph


def _hanging(paragraph, indent=Inches(0.35)) -> None:
    pf = paragraph.paragraph_format
    pf.left_indent = indent
    pf.first_line_indent = -indent


def bullet(doc, items, numbered=False, size=12):
    """A bullet or numbered list. An item may be a string, or a
    [lead, rest] pair whose lead is set in bold."""
    for i, item in enumerate(items, start=1):
        marker = f"{i}." if numbered else "•"
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        pf = paragraph.paragraph_format
        pf.space_after = Pt(4)
        pf.line_spacing = LINE_SPACING
        _hanging(paragraph, Inches(0.4))

        run = paragraph.add_run(f"{marker}\t")
        run.font.size = Pt(size)
        run.font.name = BODY_FONT

        if isinstance(item, (list, tuple)):
            lead, rest = (list(item) + [""])[:2]
            run = paragraph.add_run(lead)
            run.bold = True
            run.font.size = Pt(size)
            run.font.name = BODY_FONT
            if rest:
                run = paragraph.add_run(rest)
                run.font.size = Pt(size)
                run.font.name = BODY_FONT
        else:
            run = paragraph.add_run(item)
            run.font.size = Pt(size)
            run.font.name = BODY_FONT


def reference_list(doc, entries, size=12):
    """The bibliography, in Chicago author-date style: alphabetical by the
    first author's surname, each entry with a hanging indent and no number.

    An entry is a list of (text, italic) runs, so that Chicago's italicised
    journal and book titles are set correctly."""
    for entry in entries:
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        pf = paragraph.paragraph_format
        pf.space_after = Pt(6)
        pf.line_spacing = LINE_SPACING
        pf.left_indent = Inches(0.5)
        pf.first_line_indent = Inches(-0.5)
        parts = entry if isinstance(entry, list) else [(entry, False)]
        for text, italic in parts:
            if not text:
                continue
            run = paragraph.add_run(text)
            run.italic = italic
            run.font.size = Pt(size)
            run.font.name = BODY_FONT


def page_break(doc) -> None:
    doc.add_page_break()


# ── headings ─────────────────────────────────────────────────────────────────
def chapter_heading(doc, label: str, title: str, key: str = ""):
    """The chapter title alone, 14 pt bold capitals and centred.

    The template opens each chapter with the title only; the chapter number
    appears in the running header ("Chapter 01") and in the contents list
    ("1.  INTRODUCTION"), which is how the sample document is laid out."""
    heading = para(doc, title.upper(), size=14, bold=True, align="center",
                   space_after=18, space_before=0)
    keep_with_next(heading)
    if key:
        bookmark(heading, key)
    return heading


def front_heading(doc, title: str, key: str = "", space_after: int = 18,
                  size: int = 14):
    """A centred bold capital heading for a front-matter page.

    14 pt by default, as the template specifies for DEDICATION,
    ACKNOWLEDGMENTS and the rest; the CERTIFICATION page is the one exception
    the template sets at 12 pt."""
    paragraph = para(doc, title.upper(), size=size, bold=True, align="center",
                     space_after=space_after, space_before=0)
    keep_with_next(paragraph)
    if key:
        bookmark(paragraph, key)
    return paragraph


def heading2(doc, text: str, key: str = ""):
    """Sub-heading: ALL CAPITALS, not bold (template red note)."""
    number, _, title = text.partition("  ")
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf = paragraph.paragraph_format
    pf.space_before = Pt(12)
    pf.space_after = Pt(6)
    pf.line_spacing = LINE_SPACING
    pf.tab_stops.add_tab_stop(Inches(0.6))

    run = paragraph.add_run(f"{number}\t{title.upper()}"
                            if title else text.upper())
    run.bold = SUBHEADING_BOLD
    run.font.size = Pt(12)
    run.font.name = BODY_FONT
    keep_with_next(paragraph)
    if key:
        bookmark(paragraph, key)
    return paragraph


def heading3(doc, text: str, key: str = ""):
    """Sub-sub-heading: initial capital only, bold."""
    number, _, title = text.partition("  ")
    if title:
        title = title[:1].upper() + title[1:]
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf = paragraph.paragraph_format
    pf.space_before = Pt(10)
    pf.space_after = Pt(4)
    pf.line_spacing = LINE_SPACING
    pf.tab_stops.add_tab_stop(Inches(0.75))

    run = paragraph.add_run(f"{number}\t{title}" if title else text)
    run.bold = True
    run.font.size = Pt(12)
    run.font.name = BODY_FONT
    keep_with_next(paragraph)
    if key:
        bookmark(paragraph, key)
    return paragraph


# ── numbering counters ───────────────────────────────────────────────────────
class Counter:
    """Per-chapter figure, table and listing numbering."""

    def __init__(self):
        self.chapter = "1"
        self.fig = 0
        self.tab = 0
        self.lst = 0

    def new_chapter(self, label: str) -> None:
        self.chapter = label
        self.fig = 0
        self.tab = 0
        self.lst = 0

    def next_fig(self) -> str:
        self.fig += 1
        return f"{self.chapter}.{self.fig}"

    def next_tab(self) -> str:
        self.tab += 1
        return f"{self.chapter}.{self.tab}"

    def next_lst(self) -> str:
        self.lst += 1
        return f"{self.chapter}.{self.lst}"


# ── figures, tables, code, notes ─────────────────────────────────────────────
def figure(doc, path: str, caption: str, number: str, width: float = 5.5,
           key: str = ""):
    full = path if os.path.isabs(path) else os.path.join(FIG_DIR, path)
    holder = doc.add_paragraph()
    holder.alignment = WD_ALIGN_PARAGRAPH.CENTER
    holder.paragraph_format.space_before = Pt(8)
    holder.paragraph_format.space_after = Pt(4)
    holder.paragraph_format.line_spacing = 1.0
    if os.path.exists(full):
        holder.add_run().add_picture(full, width=Inches(width))
    else:
        run = holder.add_run(f"[missing figure: {path}]")
        run.font.color.rgb = GREY
        run.italic = True
    keep_with_next(holder)

    cap = para(doc, f"Figure {number}: {caption}", size=11, bold=True,
               align="center", space_after=12, line_spacing=1.0)
    if key:
        bookmark(cap, key)
    return cap


def placeholder_figure(doc, caption: str, number: str, guidance: str,
                       height_lines: int = 8, key: str = ""):
    """A ruled box standing in for a screenshot that is pasted in later."""
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.rows[0].cells[0]
    shade(cell, SHOT_BG)

    first = cell.paragraphs[0]
    first.alignment = WD_ALIGN_PARAGRAPH.CENTER
    first.paragraph_format.space_before = Pt(6)
    run = first.add_run(f"[ Screenshot {number} ]")
    run.bold = True
    run.font.size = Pt(11)
    run.font.name = BODY_FONT
    run.font.color.rgb = GREY

    note = cell.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = note.add_run(guidance)
    run.italic = True
    run.font.size = Pt(9)
    run.font.name = BODY_FONT
    run.font.color.rgb = GREY

    for _ in range(max(0, height_lines - 2)):
        spacer = cell.add_paragraph()
        spacer.paragraph_format.space_after = Pt(0)
        spacer.paragraph_format.line_spacing = 1.0
        spacer.add_run("").font.size = Pt(10)

    cap = para(doc, f"Figure {number}: {caption}", size=11, bold=True,
               align="center", space_after=12, space_before=4,
               line_spacing=1.0)
    if key:
        bookmark(cap, key)
    return cap


def table(doc, caption: str, number: str, headers, rows, widths=None,
          font_size: int = 10, align_left=None, key: str = ""):
    """A captioned table. The caption sits above the table, as the template
    shows for Table 3.1."""
    cap = para(doc, f"Table {number}: {caption}", size=11, bold=True,
               align="left", space_after=4, space_before=10, line_spacing=1.0)
    keep_with_next(cap)
    if key:
        bookmark(cap, key)

    grid = doc.add_table(rows=1, cols=len(headers))
    grid.style = "Table Grid"
    grid.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_cell_margins(grid)

    header_row = grid.rows[0]
    repeat_header(header_row)
    for i, text in enumerate(headers):
        cell = header_row.cells[i]
        shade(cell, HEADER_BG)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.line_spacing = 1.0
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(str(text))
        run.bold = True
        run.font.size = Pt(font_size)
        run.font.name = BODY_FONT

    for r, row in enumerate(rows):
        cells = grid.add_row().cells
        for c, value in enumerate(row):
            if c >= len(cells):
                continue
            cell = cells[c]
            if r % 2 == 1:
                shade(cell, ZEBRA_BG)
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.line_spacing = 1.0
            paragraph.paragraph_format.space_after = Pt(0)
            run = paragraph.add_run(str(value))
            run.font.size = Pt(font_size)
            run.font.name = BODY_FONT

    # Column widths are scaled to the usable text width of this template, so a
    # table sized for a wider page cannot run into the right margin.
    usable = CONTENT_WIDTH.inches
    if widths and sum(widths) > 0:
        scale = usable / sum(widths)
        scaled = [w * scale for w in widths]
    else:
        scaled = [usable / len(headers)] * len(headers)

    grid.autofit = False
    layout = _el("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    grid._tbl.tblPr.append(layout)

    # Word honours the per-cell widths; LibreOffice lays out from w:tblGrid.
    # Both are set so the document renders identically in either.
    tbl_grid = grid._tbl.find(qn("w:tblGrid"))
    if tbl_grid is not None:
        for col in list(tbl_grid):
            tbl_grid.remove(col)
        for width in scaled:
            col = _el("w:gridCol")
            col.set(qn("w:w"), str(int(width * 1440)))
            tbl_grid.append(col)

    for row in grid.rows:
        for i, width in enumerate(scaled):
            if i < len(row.cells):
                row.cells[i].width = Inches(width)

    para(doc, "", space_after=8)
    return cap


def code_block(doc, caption: str, number: str, code: str,
               path_note: str = "", key: str = ""):
    cap = para(doc, f"Listing {number}: {caption}", size=11, bold=True,
               align="left", space_after=4, space_before=10, line_spacing=1.0)
    keep_with_next(cap)
    if key:
        bookmark(cap, key)

    holder = doc.add_table(rows=1, cols=1)
    holder.style = "Table Grid"
    cell = holder.rows[0].cells[0]
    shade(cell, CODE_BG)
    set_cell_margins(holder, 80, 80, 120, 120)

    lines = code.strip("\n").split("\n")
    for i, line in enumerate(lines):
        paragraph = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf = paragraph.paragraph_format
        pf.line_spacing = 1.0
        pf.space_after = Pt(0)
        pf.space_before = Pt(0)
        run = paragraph.add_run(line if line.strip() else " ")
        run.font.name = CODE_FONT
        run.font.size = Pt(8.5)

    if path_note:
        note = para(doc, f"Source: {path_note}", size=9, italic=True,
                    align="left", space_after=10, colour=GREY,
                    line_spacing=1.0)
        note.paragraph_format.space_before = Pt(2)
    else:
        para(doc, "", space_after=8)
    return cap


def note_box(doc, text: str, title: str = "Note") -> None:
    holder = doc.add_table(rows=1, cols=1)
    holder.style = "Table Grid"
    cell = holder.rows[0].cells[0]
    shade(cell, NOTE_BG)
    set_cell_margins(holder, 90, 90, 130, 130)

    head = cell.paragraphs[0]
    head.paragraph_format.space_after = Pt(2)
    head.paragraph_format.line_spacing = 1.0
    run = head.add_run(title)
    run.bold = True
    run.font.size = Pt(10.5)
    run.font.name = BODY_FONT

    body = cell.add_paragraph()
    body.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    body.paragraph_format.line_spacing = 1.15
    body.paragraph_format.space_after = Pt(0)
    run = body.add_run(text)
    run.font.size = Pt(10.5)
    run.font.name = BODY_FONT

    para(doc, "", space_after=8)


# ── contents lists ───────────────────────────────────────────────────────────
def contents_entry(doc, text: str, page: str, level: int = 0, bold=False,
                   caps=False):
    """One line of a contents list: text on the left, page number on the right,
    joined by a dot leader."""
    paragraph = doc.add_paragraph()
    pf = paragraph.paragraph_format
    pf.space_after = Pt(2)
    pf.line_spacing = 1.15
    pf.left_indent = Inches(0.3 * level + (0.3 if level else 0))
    pf.first_line_indent = Inches(-0.3) if level else Inches(0)
    pf.tab_stops.add_tab_stop(CONTENT_WIDTH, WD_TAB_ALIGNMENT.RIGHT,
                              leader=WD_TAB_LEADER.DOTS)

    run = paragraph.add_run(text.upper() if caps else text)
    run.bold = bold
    run.font.size = Pt(11 if level else 12)
    run.font.name = BODY_FONT

    run = paragraph.add_run(f"\t{page}")
    run.bold = bold
    run.font.size = Pt(11 if level else 12)
    run.font.name = BODY_FONT
    return paragraph


# ── block renderer ───────────────────────────────────────────────────────────
def render(doc, blocks, counters, index=None):
    """Render a list of content blocks. `index` collects the bookmark keys that
    the second pass resolves into page numbers."""
    index = index if index is not None else []

    for block in blocks:
        kind = block[0]

        if kind == "chapter":
            label, title = block[1], block[2]
            number = label.split()[-1]
            counters.new_chapter(number)
            key = f"h_ch{number}"
            chapter_heading(doc, label, title, key=key)
            index.append({"kind": "toc", "level": 0, "key": key,
                          "text": f"{number}.  {title}", "caps": True,
                          "bold": True, "anchor": title,
                          "opens_page": True})

        elif kind == "h1":
            title = block[1]
            if len(block) > 2:
                label = block[2]
            elif title.upper().startswith("APPENDIX "):
                # "APPENDIX B:  SURVEY" numbers its figures B.1, B.2, ...
                label = title.split()[1].rstrip(":")
            else:
                label = title[:1]
            counters.new_chapter(label)
            key = f"h_{label}"
            front_heading(doc, title, key=key)
            index.append({"kind": "toc", "level": 0, "key": key,
                          "text": title, "caps": True, "bold": True,
                          "anchor": title, "opens_page": True})

        elif kind == "h2":
            text = block[1]
            key = f"h2_{len(index)}"
            heading2(doc, text, key=key)
            index.append({"kind": "toc", "level": 1, "key": key,
                          "text": text, "caps": True, "bold": False})

        elif kind == "h3":
            text = block[1]
            key = f"h3_{len(index)}"
            heading3(doc, text, key=key)
            index.append({"kind": "toc", "level": 2, "key": key,
                          "text": text, "caps": False, "bold": False})

        elif kind == "p":
            para(doc, block[1], indent=True)

        elif kind == "bullets":
            bullet(doc, block[1])

        elif kind == "numbers":
            bullet(doc, block[1], numbered=True)

        elif kind == "refs":
            reference_list(doc, block[1])

        elif kind == "note":
            note_box(doc, block[1], block[2] if len(block) > 2 else "Note")

        elif kind == "fig":
            path, caption = block[1], block[2]
            width = block[3] if len(block) > 3 else 5.5
            number = counters.next_fig()
            key = f"fig_{number}"
            figure(doc, path, caption, number, width, key=key)
            index.append({"kind": "lof", "number": number, "key": key,
                          "caption": caption})

        elif kind == "shot":
            caption = block[2]
            guidance = block[3] if len(block) > 3 else ""
            number = counters.next_fig()
            key = f"fig_{number}"
            placeholder_figure(doc, caption, number, guidance, key=key)
            index.append({"kind": "lof", "number": number, "key": key,
                          "caption": caption})

        elif kind == "table":
            caption, headers, rows = block[1], block[2], block[3]
            widths = block[4] if len(block) > 4 else None
            size = block[5] if len(block) > 5 else 10
            number = counters.next_tab()
            key = f"tab_{number}"
            table(doc, caption, number, headers, rows, widths, size, key=key)
            index.append({"kind": "lot", "number": number, "key": key,
                          "caption": caption})

        elif kind == "code":
            caption, code = block[1], block[2]
            path_note = block[3] if len(block) > 3 else ""
            number = counters.next_lst()
            code_block(doc, caption, number, code, path_note)

        elif kind == "pagebreak":
            page_break(doc)

        else:
            raise ValueError(f"unknown block kind: {kind!r}")

    return index
