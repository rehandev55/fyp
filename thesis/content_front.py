"""
Front matter.

The page order follows the template: cover, inner title page, certification,
dedication, contents, list of figures, list of tables, list of abbreviations,
acknowledgments, abstract. Roman numerals run from the cover to the abstract;
Chapter 1 restarts the count in arabic numerals.
"""

from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.shared import Inches, Pt

import os

from docx_kit import (FIG_DIR, bookmark, contents_entry, front_heading,
                      page_break, para, rich)

TITLE = ("ETERNAL SUNSHINE: AN AI-POWERED BOARD-SPECIFIC LEARNING, "
         "ASSESSMENT AND RESOURCE PLATFORM FOR CLASSES 9–12")
TITLE_QUOTED = ("Eternal Sunshine: An AI-Powered Board-Specific Learning, "
                "Assessment and Resource Platform for Classes 9–12")

STUDENTS = [
    ("Anum Sheraz", "FA22-BSE-007"),
    ("Munazza Javed", "FA22-BSE-019"),
    ("Aniqa Saghir", "FA22-BSE-031"),
]
SESSION = "Session 2022–2026"
SUPERVISOR = "Engr. Areeb Ahmed Mir"
DEPT = "Department of Software Engineering"
FACULTY = "Faculty of Engineering & Technology"
UNIVERSITY = "Mirpur University of Science and Technology (MUST), Mirpur"
COUNTRY = "AJK Pakistan"

ABBREVIATIONS = [
    ("AI", "Artificial Intelligence"),
    ("AJK", "Azad Jammu and Kashmir"),
    ("API", "Application Programming Interface"),
    ("BISE", "Board of Intermediate and Secondary Education"),
    ("CRUD", "Create, Read, Update, Delete"),
    ("CSRF", "Cross-Site Request Forgery"),
    ("DFD", "Data Flow Diagram"),
    ("ERD", "Entity-Relationship Diagram"),
    ("FBISE", "Federal Board of Intermediate and Secondary Education"),
    ("HTTP / HTTPS", "Hypertext Transfer Protocol (Secure)"),
    ("JSON", "JavaScript Object Notation"),
    ("JSONL", "JSON Lines (one JSON object per line)"),
    ("LLM", "Large Language Model"),
    ("LRU", "Least Recently Used (cache eviction policy)"),
    ("MAE", "Mean Absolute Error"),
    ("MCQ", "Multiple Choice Question"),
    ("MVC", "Model-View-Controller"),
    ("OAuth", "Open Authorisation"),
    ("OCR", "Optical Character Recognition"),
    ("ORM", "Object-Relational Mapping"),
    ("RAG", "Retrieval-Augmented Generation"),
    ("REST", "Representational State Transfer"),
    ("SDG", "Sustainable Development Goal"),
    ("SPA", "Single Page Application"),
    ("SQL", "Structured Query Language"),
    ("SRS", "Software Requirements Specification"),
    ("TLS", "Transport Layer Security"),
    ("TOTP", "Time-based One-Time Password"),
    ("UAT", "User Acceptance Testing"),
    ("UI / UX", "User Interface / User Experience"),
    ("UML", "Unified Modeling Language"),
]


def _blank(doc, n=1):
    for _ in range(n):
        para(doc, "", space_after=0)


def _c(doc, text, size=14, bold=True, space_after=0):
    return para(doc, text, size=size, bold=bold, align="center",
                space_after=space_after, line_spacing=1.0)


# ── page 1: cover ────────────────────────────────────────────────────────────
def cover_page(doc):
    _blank(doc, 1)
    _c(doc, TITLE, size=16, space_after=0)
    _blank(doc, 2)

    logo = os.path.join(FIG_DIR, "logo_must.jpg")
    if os.path.exists(logo):
        holder = para(doc, "", align="center", space_after=0, line_spacing=1.0)
        holder.add_run().add_picture(logo, width=Inches(1.5))
    _blank(doc, 2)

    for name, reg in STUDENTS:
        _c(doc, name, size=14)
        _c(doc, f"({reg})", size=14, space_after=6)
    _blank(doc, 3)

    _c(doc, SESSION, size=14, space_after=0)
    _blank(doc, 6)

    _c(doc, DEPT, size=14)
    _c(doc, FACULTY, size=14)
    _c(doc, UNIVERSITY, size=14)
    _c(doc, COUNTRY, size=14)
    page_break(doc)


# ── page 2: inner title ──────────────────────────────────────────────────────
def inner_title_page(doc):
    _blank(doc, 1)
    _c(doc, TITLE, size=16, space_after=0)
    _blank(doc, 1)

    logo = os.path.join(FIG_DIR, "logo_must.jpg")
    if os.path.exists(logo):
        holder = para(doc, "", align="center", space_after=0, line_spacing=1.0)
        holder.add_run().add_picture(logo, width=Inches(1.2))
    _blank(doc, 1)

    _c(doc, "By", size=14, bold=False, space_after=6)
    for name, reg in STUDENTS:
        _c(doc, name, size=14)
        _c(doc, f"({reg})", size=14, space_after=6)
    _blank(doc, 2)

    _c(doc, "A Project Report is submitted in partial fulfillment of",
       size=14, bold=False)
    _c(doc, "the requirement for the degree of", size=14, bold=False,
       space_after=6)
    _c(doc, "Bachelor of Science", size=14)
    _c(doc, "In", size=14, bold=False)
    _c(doc, "SOFTWARE ENGINEERING", size=14, space_after=6)
    _blank(doc, 2)

    _c(doc, SESSION, size=14, bold=False, space_after=0)
    _blank(doc, 4)

    _c(doc, DEPT, size=14)
    _c(doc, FACULTY, size=14)
    _c(doc, f"{UNIVERSITY} (AJK) Pakistan", size=14)
    page_break(doc)


# ── page 3: certification ────────────────────────────────────────────────────
def certification_page(doc):
    front_heading(doc, "CERTIFICATION", key="fm_certification",
                  size=12)

    para(doc,
         "We hereby undertake that this final year project is an original one "
         "and no part of this final year project falls under plagiarism. If "
         "found otherwise, at any stage, we will be responsible for the "
         "consequences.",
         indent=True, space_after=14)

    for name, reg in STUDENTS:
        line = para(doc, "", space_after=4, line_spacing=1.0)
        line.paragraph_format.tab_stops.add_tab_stop(Inches(3.2))
        run = line.add_run(f"Student’s Name: {name}")
        run.font.size = Pt(12)
        run = line.add_run(f"\tSignature: __________")
        run.font.size = Pt(12)

        line = para(doc, "", space_after=12, line_spacing=1.0)
        line.paragraph_format.tab_stops.add_tab_stop(Inches(3.2))
        run = line.add_run(f"Registration No.: {reg}")
        run.font.size = Pt(12)
        run = line.add_run("\tDate: ______________")
        run.font.size = Pt(12)

    para(doc, "Certified that the contents and form of final year project "
              "entitled", align="center", space_after=6, space_before=8)
    para(doc, f"“{TITLE_QUOTED}”", align="center", bold=True,
         space_after=6)
    para(doc, "submitted by", align="center", space_after=6)
    para(doc, ", ".join(name for name, _ in STUDENTS), align="center",
         bold=True, space_after=6)
    para(doc, "have been found satisfactory for the requirement of the "
              "degree.", align="center", space_after=18)

    for label in ("Supervisor", "Co-Supervisor (If any)",
                  "External Examiner"):
        para(doc, f"{label}: ___________________________________",
             align="right", space_after=10, line_spacing=1.0)
    para(doc, "Chairperson: ___________________________________",
         align="right", space_after=0, line_spacing=1.0)
    page_break(doc)


# ── page 4: dedication ───────────────────────────────────────────────────────
def dedication_page(doc):
    front_heading(doc, "DEDICATION", key="fm_dedication")
    _blank(doc, 4)
    para(doc,
         "To our parents, whose patience and quiet sacrifice made every late "
         "night possible; to our teachers in the Department of Software "
         "Engineering, who taught us to ask better questions before writing a "
         "single line of code; and to the students of classes nine to twelve "
         "across Pakistan who study without a tutor, a library or a second "
         "chance — this work is for you.",
         align="center", space_after=0)
    page_break(doc)


# ── contents, list of figures, list of tables ────────────────────────────────
#: Front-matter pages that the template lists inside CONTENTS, with roman
#: page numbers. The key is the bookmark the build resolves the page from.
FRONT_ENTRIES = [
    ("fm_lof", "LIST OF FIGURES"),
    ("fm_lot", "LIST OF TABLES"),
    ("fm_abbrev", "LIST OF ABBREVIATIONS"),
    ("fm_ack", "ACKNOWLEDGMENTS"),
    ("fm_abstract", "ABSTRACT"),
]


def contents_page(doc, entries, front=None):
    """`entries` is the resolved table of contents: a list of dicts with
    level, text, caps, bold and page. `front` maps a front-matter bookmark to
    its roman page number. Both are empty on the first build pass."""
    front_heading(doc, "CONTENTS", key="fm_contents")
    head = para(doc, "", align="right", space_after=6, line_spacing=1.0)
    run = head.add_run("Page No.")
    run.bold = True
    run.font.size = Pt(12)

    if not entries:
        para(doc, "[contents resolved on the second build pass]",
             align="center", space_after=0)

    front = front or {}
    for key, label in FRONT_ENTRIES:
        if key in front:
            contents_entry(doc, label, front[key], level=0, bold=True,
                           caps=True)

    for entry in entries:
        contents_entry(doc, entry["text"], entry["page"],
                       level=entry["level"], bold=entry["bold"],
                       caps=entry["caps"])
    page_break(doc)


def list_of_figures(doc, entries):
    front_heading(doc, "LIST OF FIGURES", key="fm_lof")
    head = para(doc, "", space_after=6, line_spacing=1.0)
    head.paragraph_format.tab_stops.add_tab_stop(
        Inches(5.59), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
    run = head.add_run("Figure No.")
    run.bold = True
    run.font.size = Pt(12)
    run = head.add_run("\tPage No.")
    run.bold = True
    run.font.size = Pt(12)

    if not entries:
        para(doc, "[list resolved on the second build pass]", align="center")
    for entry in entries:
        contents_entry(doc, f"{entry['number']}: {entry['caption']}",
                       entry["page"], level=0)
    page_break(doc)


def list_of_tables(doc, entries):
    front_heading(doc, "LIST OF TABLES", key="fm_lot")
    head = para(doc, "", space_after=6, line_spacing=1.0)
    head.paragraph_format.tab_stops.add_tab_stop(
        Inches(5.59), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
    run = head.add_run("Table No.")
    run.bold = True
    run.font.size = Pt(12)
    run = head.add_run("\tPage No.")
    run.bold = True
    run.font.size = Pt(12)

    if not entries:
        para(doc, "[list resolved on the second build pass]", align="center")
    for entry in entries:
        contents_entry(doc, f"{entry['number']}  {entry['caption']}",
                       entry["page"], level=0)
    page_break(doc)


def abbreviations_page(doc):
    front_heading(doc, "LIST OF ABBREVIATIONS", key="fm_abbrev")
    for short, long in ABBREVIATIONS:
        line = para(doc, "", space_after=2, line_spacing=1.15)
        line.paragraph_format.tab_stops.add_tab_stop(Inches(1.4))
        line.paragraph_format.left_indent = Inches(1.4)
        line.paragraph_format.first_line_indent = Inches(-1.4)
        run = line.add_run(f"{short}\t")
        run.bold = True
        run.font.size = Pt(12)
        run = line.add_run(long)
        run.font.size = Pt(12)
    page_break(doc)


# ── acknowledgments ──────────────────────────────────────────────────────────
def acknowledgments_page(doc):
    front_heading(doc, "ACKNOWLEDGMENTS", key="fm_ack")

    for text in [
        "All praise is due to Almighty Allah, who gave us the health, the "
        "patience and the clear mind needed to finish this work.",

        "We are deeply grateful to our supervisor, Engr. Areeb Ahmed Mir, of "
        "the Department of Software Engineering, for his guidance through "
        "both phases of this project. He insisted that we test our "
        "assumptions on real students before fixing a design. That single "
        "instruction changed the direction of the project for the better. His "
        "review of our architecture also saved us from several costly "
        "mistakes.",

        "We thank the Chairperson and the faculty of the Department of "
        "Software Engineering for the academic environment, the laboratory "
        "access and the review feedback that shaped this report. The comments "
        "we received at the proposal defence and at the Software Requirements "
        "Specification review were carried directly into the requirements and "
        "the architecture presented here.",

        "We thank the forty-eight students of classes nine to twelve who "
        "answered our requirement survey, and the volunteers who took part in "
        "the usability sessions. They told us what they actually struggle "
        "with, rather than what we assumed they struggle with. That is the "
        "reason the platform looks the way it does.",

        "Finally, we thank our families and friends for their encouragement "
        "and their patience during the long months of development, testing "
        "and writing.",
    ]:
        para(doc, text, indent=True, space_after=8)

    _blank(doc, 2)
    for name, _ in STUDENTS:
        para(doc, name, align="right", bold=True, space_after=2,
             line_spacing=1.0)
    page_break(doc)


# ── abstract: one page ───────────────────────────────────────────────────────
def abstract_page(doc):
    front_heading(doc, "ABSTRACT", key="fm_abstract", space_after=10)

    for text in [
        "School education in Pakistan is organised around examination boards, "
        "each with its own syllabus, prescribed textbooks and papers. The "
        "digital help available to students is not organised that way: a "
        "search engine, a video site or a general chatbot answers every "
        "student alike, whatever their board. In a survey of forty-eight "
        "students run for this project, 79.2 per cent said online material "
        "differs from what their board examines, and 52.1 per cent said "
        "past papers are hard to find.",

        "This report presents Eternal Sunshine, a web platform that ties "
        "learning, testing and study resources to a named board, class and "
        "subject. It has three tiers: a React interface, a Laravel server and "
        "a separate FastAPI service for the artificial-intelligence work. "
        "Its core is Retrieval-Augmented Generation. Board textbooks, "
        "keybooks and past papers are converted to text by optical character "
        "recognition, cleaned, chunked and stored in a vector index of 8,057 "
        "passages, and every request to the model is grounded in passages "
        "fetched under a strict filter on board, class and subject. Three "
        "agents share it: a teacher that explains, a tester that sets "
        "labelled questions, and an evaluator that marks written answers.",

        "Thirty curriculum questions all returned relevant material, and all "
        "sixty checked passages matched the requested board, class and "
        "subject, a filter precision of 1.00. Judged by an independent model "
        "with retrieval on and off, grounding rose from 1.92 to 2.75 out of "
        "5, and retrieval never scored worse. The evaluator matched the "
        "expected mark exactly on 73.3 per cent of a fifteen-answer gold set "
        "and was within one mark on every item. Generated question sets "
        "passed all twenty-one structural checks, and the memory manager held "
        "prompt size flat, saving 37 per cent of tokens by turn ten. "
        "Built-in accounting places a tutoring exchange well under one US "
        "cent.",

        "Grounding a small, inexpensive model in a curated, board-filtered "
        "corpus is therefore a practical route to trustworthy educational "
        "artificial intelligence in a low-resource setting. Wider board "
        "coverage, stricter grounding and a teacher dashboard are the main "
        "future work.",
    ]:
        para(doc, text, indent=True, space_after=0)

    keywords = para(doc, "", indent=False, space_after=0, space_before=6)
    run = keywords.add_run("Keywords: ")
    run.bold = True
    run.font.size = Pt(12)
    run = keywords.add_run(
        "educational technology; Retrieval-Augmented Generation; large "
        "language models; automated assessment; Pakistani board curriculum.")
    run.font.size = Pt(12)


def build(doc, toc=None, lof=None, lot=None, front=None):
    cover_page(doc)
    inner_title_page(doc)
    certification_page(doc)
    dedication_page(doc)
    contents_page(doc, toc or [], front or {})
    list_of_figures(doc, lof or [])
    list_of_tables(doc, lot or [])
    abbreviations_page(doc)
    acknowledgments_page(doc)
    abstract_page(doc)
